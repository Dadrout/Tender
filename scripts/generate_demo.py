"""Generated synthetic source documents, parsed through the regular pipeline."""

from pathlib import Path
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    PageBreak,
    Table,
    TableStyle,
)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "demo"
OUT.mkdir(exist_ok=True)
pdfmetrics.registerFont(
    TTFont("DemoSans", str(ROOT / "apps/api/qazbid/assets/NotoSans-Regular.ttf"))
)
style = ParagraphStyle(
    "body", fontName="DemoSans", fontSize=11, leading=19, spaceAfter=9
)


def pdf(slug, title, product, qty, attrs, budget, days):
    story = [
        Paragraph("QazBid AI / СИНТЕТИЧЕСКИЙ УЧЕБНЫЙ ПРИМЕР", style),
        Spacer(1, 12),
    ]
    lines = [
        f"Название: {title}",
        f"Номер: DEMO-{slug.upper()}-2026",
        "Заказчик: Учебная муниципальная организация",
        "Регион: Тараз, Жамбылская область",
        "Способ закупки: Учебный запрос ценовых предложений",
        f"Бюджет: {budget} KZT",
        "Срок подачи: 30.10.2026",
        f"Товар: {product}",
        f"Количество: {qty}",
        "Единица: шт",
    ] + [f"{k}: {v}" for k, v in attrs.items()]
    for line in lines:
        story.append(Paragraph(line, style))
    table = Table(
        [["Позиция", "Количество", "Единица"], [product, str(qty), "шт"]],
        colWidths=[320, 90, 70],
    )
    table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), "DemoSans"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eaf0fb")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#ccd6e8")),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
            ]
        )
    )
    story.append(table)
    story.append(PageBreak())
    for line in [
        "ТЕХНИЧЕСКИЕ И ДОГОВОРНЫЕ УСЛОВИЯ",
        "Документ: Сертификат соответствия",
        f"Поставка: в течение {days} дней, до склада заказчика в Таразе",
        "Оплата: в течение 30 дней после приёмки",
        "Требуется поставка нового товара в оригинальной упаковке.",
        "Учебные условия не являются действующей закупкой.",
    ]:
        story.append(Paragraph(line, style))
    SimpleDocTemplate(
        str(OUT / f"tender-{slug}.pdf"), pagesize=A4, rightMargin=42, leftMargin=42
    ).build(story)


pdf(
    "enamel",
    "Поставка эмали для хозяйственных работ",
    "Эмаль жёлтая ПФ-115",
    12,
    {"Объем": "3.5 л", "Цвет": "Жёлтый", "Стандарт": "ГОСТ 6465-76"},
    90000,
    5,
)
pdf(
    "brushes",
    "Поставка малярных инструментов",
    "Кисть малярная плоская",
    30,
    {"Размер": "50 мм", "Материал": "Натуральная щетина"},
    60000,
    7,
)
pdf(
    "buckets",
    "Поставка хозяйственного инвентаря",
    "Ведро пластиковое",
    20,
    {"Объем": "10 л", "Материал": "Пластик"},
    50000,
    3,
)
headers = [
    "Наименование",
    "Артикул",
    "Количество",
    "Единица измерения",
    "Закупочная цена",
    "Цена поставщика",
    "Объем",
    "Размер",
    "Цвет",
    "Материал",
    "Стандарт",
    "Сертификат",
    "Поставщик",
    "Срок поставки",
    "Наличие поставщика",
]
rows = [
    [
        "Эмаль жёлтая ПФ-115",
        "EN-35-Y",
        8,
        "шт",
        4000,
        4500,
        "3.5 л",
        "",
        "Жёлтый",
        "",
        "ГОСТ 6465-76",
        "",
        "Тараз Снаб",
        4,
        "Да",
    ],
    [
        "Эмаль жёлтая ПФ-115",
        "EN-25-Y",
        25,
        "шт",
        3100,
        3400,
        "2.5 л",
        "",
        "Жёлтый",
        "",
        "ГОСТ 6465-76",
        "",
        "Тараз Снаб",
        4,
        "Да",
    ],
    [
        "Эмаль белая ПФ-115",
        "EN-35-W",
        18,
        "шт",
        3900,
        4200,
        "3.5 л",
        "",
        "Белый",
        "",
        "ГОСТ 6465-76",
        "",
        "Тараз Снаб",
        3,
        "Да",
    ],
    [
        "Эмаль жёлтая универсальная",
        "EN-UNK",
        6,
        "шт",
        2800,
        "",
        "",
        "",
        "Жёлтый",
        "",
        "",
        "",
        "Местный склад",
        "",
        "Нет",
    ],
    [
        "Кисть малярная плоская",
        "BR-50",
        40,
        "шт",
        800,
        1000,
        "",
        "50 мм",
        "",
        "Натуральная щетина",
        "",
        "",
        "Мастер Снаб",
        2,
        "Да",
    ],
    [
        "Кисть малярная плоская",
        "BR-40",
        15,
        "шт",
        600,
        750,
        "",
        "40 мм",
        "",
        "Синтетика",
        "",
        "",
        "Мастер Снаб",
        2,
        "Да",
    ],
    [
        "Ведро пластиковое",
        "BU-10",
        10,
        "шт",
        950,
        1200,
        "10 л",
        "",
        "",
        "Пластик",
        "",
        "",
        "Тараз Снаб",
        6,
        "Да",
    ],
    [
        "Ведро пластиковое",
        "BU-12",
        22,
        "шт",
        1100,
        1300,
        "12 л",
        "",
        "",
        "Пластик",
        "",
        "",
        "Тараз Снаб",
        2,
        "Да",
    ],
    [
        "Валик малярный",
        "RL-20",
        20,
        "шт",
        1500,
        1750,
        "",
        "200 мм",
        "",
        "Полиамид",
        "",
        "",
        "Мастер Снаб",
        2,
        "Да",
    ],
    [
        "Перчатки хозяйственные",
        "GL-L",
        50,
        "шт",
        350,
        420,
        "",
        "L",
        "",
        "Латекс",
        "",
        "",
        "Жамбыл Опт",
        1,
        "Да",
    ],
    [
        "Мыло хозяйственное",
        "SOAP",
        70,
        "шт",
        200,
        240,
        "",
        "",
        "",
        "",
        "",
        "",
        "Жамбыл Опт",
        1,
        "Да",
    ],
    [
        "Швабра с насадкой",
        "MOP",
        12,
        "шт",
        2100,
        2500,
        "",
        "",
        "",
        "Пластик",
        "",
        "",
        "Жамбыл Опт",
        3,
        "Да",
    ],
    [
        "Лампа LED 12 Вт",
        "LED12",
        35,
        "шт",
        600,
        700,
        "",
        "",
        "Белый",
        "",
        "",
        "",
        "ЭлектроТараз",
        2,
        "Да",
    ],
    [
        "Бумага офисная А4",
        "PAPER",
        30,
        "шт",
        1700,
        1900,
        "",
        "210x297 мм",
        "",
        "",
        "",
        "",
        "Офис Снаб",
        2,
        "Да",
    ],
    [
        "Пакеты для мусора",
        "BAG60",
        80,
        "шт",
        400,
        480,
        "60 л",
        "",
        "Чёрный",
        "Полиэтилен",
        "",
        "",
        "Жамбыл Опт",
        1,
        "Да",
    ],
    [
        "Цемент М400 50 кг",
        "CEMENT",
        45,
        "шт",
        1800,
        2000,
        "",
        "",
        "",
        "",
        "",
        "",
        "Строй Снаб",
        2,
        "Да",
    ],
    [
        "Скотч упаковочный",
        "TAPE",
        60,
        "шт",
        300,
        370,
        "",
        "50 мм",
        "",
        "",
        "",
        "",
        "Офис Снаб",
        2,
        "Да",
    ],
]
wb = Workbook()
ws = wb.active
ws.title = "Каталог"
ws.append(headers)
for row in rows:
    ws.append(row)
ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions
for cell in ws[1]:
    cell.fill = PatternFill("solid", fgColor="203650")
    cell.font = Font(color="FFFFFF", bold=True)
    cell.alignment = Alignment(wrap_text=True)
for col in ws.columns:
    ws.column_dimensions[col[0].column_letter].width = 32 if col[0].column == 1 else 21
ws.row_dimensions[1].height = 38
wb.save(OUT / "inventory.xlsx")
print("Generated 3 two-page synthetic PDFs and 17-row inventory.xlsx")
